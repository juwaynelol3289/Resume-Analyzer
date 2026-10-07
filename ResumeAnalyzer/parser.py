from pathlib import Path


class ResumeParserError(Exception):
    """Custom error for resume parsing problems."""
    pass


def parse_txt(file_path: str) -> str:
    """Extract text from a TXT file."""
    try:
        return Path(file_path).read_text(encoding="utf-8")
    except UnicodeDecodeError:
        raise ResumeParserError("The TXT file is not encoded as UTF-8.")
    except OSError as e:
        raise ResumeParserError(f"Could not read TXT file: {e}")


def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    try:
        import PyPDF2
    except ImportError:
        raise ResumeParserError(
            "PyPDF2 is not installed. Run: pip install PyPDF2"
        )

    try:
        text = []

        with open(file_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)

            for page in reader.pages:
                page_text = page.extract_text()

                if page_text:
                    text.append(page_text)

        return "\n".join(text)

    except OSError as e:
        raise ResumeParserError(f"Could not open PDF: {e}")
    except Exception as e:
        raise ResumeParserError(f"Could not parse PDF: {e}")


def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    try:
        from docx import Document
    except ImportError:
        raise ResumeParserError(
            "python-docx is not installed. Run: pip install python-docx"
        )

    try:
        document = Document(file_path)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Also extract text from tables.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        return "\n".join(paragraphs)

    except OSError as e:
        raise ResumeParserError(f"Could not open DOCX: {e}")
    except Exception as e:
        raise ResumeParserError(f"Could not parse DOCX: {e}")


def parse_resume(file_path: str) -> str:
    """
    Automatically detect the resume format and extract its text.

    Supported formats:
        .txt
        .pdf
        .docx
    """
    path = Path(file_path)

    if not path.exists():
        raise ResumeParserError(f"File does not exist: {file_path}")

    if not path.is_file():
        raise ResumeParserError(f"Not a file: {file_path}")

    extension = path.suffix.lower()

    if extension == ".txt":
        text = parse_txt(file_path)

    elif extension == ".pdf":
        text = parse_pdf(file_path)

    elif extension == ".docx":
        text = parse_docx(file_path)

    else:
        raise ResumeParserError(
            f"Unsupported file type: {extension}\n"
            "Supported formats: .txt, .pdf, .docx"
        )

    # Clean up excessive blank lines.
    cleaned_lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    return "\n".join(cleaned_lines)


if __name__ == "__main__":
    # Simple test
    file = input("Enter resume file path: ").strip()

    try:
        resume_text = parse_resume(file)

        print("\n" + "=" * 50)
        print("RESUME TEXT")
        print("=" * 50)
        print(resume_text)

    except ResumeParserError as e:
        print(f"\n❌ Error: {e}")